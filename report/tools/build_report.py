import html
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from content import BLOCKS

ROOT = Path(__file__).resolve().parents[1]
REPO = ROOT.parent
ASSETS = ROOT / "assets"

FONT_BODY = "Times New Roman"
FONT_MONO = "Consolas"
CONTENT_WIDTH_IN = 6.5


def slugify(text: str) -> str:
    value = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return value or "section"


def collect() -> tuple[list, list, list]:
    toc = []
    figures = []
    tables = []
    figure_no = 0
    table_no = 0
    for block in BLOCKS:
        kind = block[0]
        if kind in ("h1", "h2"):
            toc.append((1 if kind == "h1" else 2, block[1], slugify(block[1])))
        elif kind == "figure":
            figure_no += 1
            data = block[1]
            figures.append({
                "no": figure_no,
                "caption": data["caption"],
                "key": f"fig-{figure_no}",
                "file": data["file"],
                "width": data.get("width", 6.2),
            })
        elif kind == "table":
            table_no += 1
            tables.append({
                "no": table_no,
                "caption": block[1]["caption"],
                "key": f"tab-{table_no}",
            })
    return toc, figures, tables


TOC, FIGURES, TABLES = collect()


def page_no(page_numbers: dict, key: str) -> str:
    value = page_numbers.get(key)
    return str(value) if value else ""


def inline_html(text: str) -> str:
    value = html.escape(text, quote=False)
    value = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", value)
    value = re.sub(r"`([^`]+)`", r"<code>\1</code>", value)
    value = re.sub(r"(https?://[^\s<]+)", r'<a href="\1">\1</a>', value)
    return value


def html_page(page_numbers: dict) -> str:
    parts = []
    parts.append(
        "<!doctype html><html lang='en'><head><meta charset='utf-8'>"
        "<title>RailReserve - Preventing Double Booking in a Railway Reservation System</title>"
        "<style>"
        "@page { size: A4; margin: 24mm 20mm 22mm 20mm; }"
        "html { -webkit-print-color-adjust: exact; print-color-adjust: exact; }"
        "body { margin:0; color:#111827; font-family:'Liberation Serif','Times New Roman',serif; font-size:11pt; line-height:1.38; }"
        "p { margin:0 0 6pt; text-align:justify; orphans:3; widows:3; }"
        "h1 { font-size:14pt; font-weight:bold; margin:0 0 6pt; page-break-after:avoid; }"
        "h2 { font-size:12pt; font-weight:bold; margin:10pt 0 4pt; page-break-after:avoid; }"
        ".front { font-size:14pt; font-weight:bold; margin:0 0 6pt; page-break-after:avoid; }"
        "strong { font-weight:bold; }"
        "code { font-family:'DejaVu Sans Mono',Consolas,monospace; font-size:9.2pt; }"
        "a { color:#0b4f6c; text-decoration:none; }"
        "ul,ol { margin:0 0 6pt 0; padding-left:18pt; }"
        "li { margin:0 0 2.5pt; text-align:justify; }"
        ".pagebreak { page-break-before:always; }"
        ".title-page { height:249mm; box-sizing:border-box; padding:10mm 12mm 6mm; text-align:center; page-break-after:always; overflow:hidden; }"
        ".title-page p { text-align:center; margin:0; }"
        ".trust { font-size:10.4pt; line-height:1.35; }"
        ".college { font-size:11.2pt; font-weight:bold; line-height:1.3; }"
        ".address,.autonomy,.affiliation { font-size:10.2pt; line-height:1.32; }"
        ".dept { font-size:11.4pt; font-weight:bold; margin-top:20mm !important; letter-spacing:.2pt; }"
        ".year { font-size:10.8pt; margin-top:1.5mm !important; }"
        ".logo { width:27mm; margin:9mm auto 6mm; display:block; }"
        ".report-label { font-size:10.6pt; margin-top:1mm !important; }"
        ".report-title { font-size:18.5pt; font-weight:bold; line-height:1.22; margin:7mm 12mm 4mm !important; }"
        ".subtitle { font-size:10.8pt; font-style:italic; margin-bottom:11mm !important; }"
        ".by { font-size:10.6pt; margin-bottom:6mm !important; }"
        ".student { font-size:11pt; line-height:1.6; }"
        ".degree { font-size:10.8pt; font-weight:bold; margin-top:11mm !important; }"
        ".toc { margin-top:3pt; }"
        ".toc-row { display:flex; align-items:baseline; font-size:9.3pt; line-height:1.3; margin:0; }"
        ".toc-row .t { flex:0 1 auto; }"
        ".toc-row .dots { flex:1 1 auto; border-bottom:1px dotted #9aa3ad; margin:0 4pt 2pt; min-width:10pt; }"
        ".toc-row .pg { flex:0 0 24pt; text-align:right; }"
        ".toc-row.lvl2 { padding-left:7mm; font-size:9.0pt; }"
        ".list-table { width:100%; border-collapse:collapse; font-size:9.1pt; line-height:1.28; margin:0 0 6pt; }"
        ".list-table td { border:none; padding:0.8pt 0; vertical-align:top; }"
        ".list-table td.pg { text-align:right; width:22pt; }"
        ".abbr { width:100%; border-collapse:collapse; font-size:9.2pt; margin-top:2pt; }"
        ".abbr td { border:0.5pt solid #b9c2cc; padding:2.6pt 4pt; }"
        ".table-block { margin:6pt 0 8pt; page-break-inside:avoid; }"
        ".table-caption { font-size:9.4pt; font-style:italic; margin:0 0 3pt; }"
        "table.data { width:100%; border-collapse:collapse; font-size:8.8pt; line-height:1.26; }"
        "table.data th { background:#dce6f1; border:0.55pt solid #8595a6; padding:3.2pt 4pt; text-align:left; vertical-align:top; font-weight:bold; }"
        "table.data td { border:0.55pt solid #9caaba; padding:3.2pt 4pt; vertical-align:top; }"
        "table.data tr { page-break-inside:avoid; }"
        ".figure-block { margin:7pt 0 9pt; page-break-inside:avoid; text-align:center; }"
        ".figure-block img { max-width:100%; display:block; margin:0 auto; }"
        ".figure-caption { font-size:9.4pt; font-style:italic; margin-top:3pt; text-align:center; }"
        ".code-block { margin:5pt 0 8pt; page-break-inside:avoid; }"
        ".code-caption { font-size:9.4pt; font-style:italic; margin:0 0 3pt; }"
        ".code-block pre { margin:0; padding:6pt 8pt; background:#f5f7fa; border-left:2.6pt solid #5b6b7c; border-top:0.5pt solid #d3d9e0; border-right:0.5pt solid #d3d9e0; border-bottom:0.5pt solid #d3d9e0; font-family:'DejaVu Sans Mono',Consolas,monospace; font-size:8.2pt; line-height:1.3; white-space:pre-wrap; }"
        ".ref { font-size:9.3pt; line-height:1.3; margin:0 0 4pt; padding-left:8mm; text-indent:-8mm; text-align:left; }"
        ".ref a { word-break:break-all; }"
        ".qa { margin:0 0 6pt; }"
        ".qa .q { font-weight:bold; }"
        ".small { font-size:9.5pt; }"
        "</style></head><body>"
    )

    for block in BLOCKS:
        kind = block[0]
        if kind == "title":
            data = block[1]
            parts.append("<section class='title-page'>")
            parts.append(f"<p class='trust'>{html.escape(data['trust'])}</p>")
            parts.append(f"<p class='college'>{html.escape(data['college'])}</p>")
            parts.append(f"<p class='address'>{html.escape(data['address'])}</p>")
            parts.append(f"<p class='autonomy'>{html.escape(data['autonomy'])}</p>")
            parts.append(f"<p class='affiliation'>{html.escape(data['affiliation'])}</p>")
            parts.append(f"<p class='dept'>{html.escape(data['department'])}</p>")
            parts.append(f"<p class='year'>{html.escape(data['year'])}</p>")
            parts.append("<img class='logo' src='assets/pccoe-logo.png' alt='PCCoE logo'>")
            parts.append(f"<p class='report-label'>{html.escape(data['report'])}</p>")
            parts.append(f"<p class='report-title'>{html.escape(data['title'])}</p>")
            parts.append(f"<p class='subtitle'>{html.escape(data['subtitle'])}</p>")
            parts.append(f"<p class='by'>{html.escape(data['by'])}</p>")
            for name, roll in data["students"]:
                parts.append(f"<p class='student'>{html.escape(name)}: {html.escape(roll)}</p>")
            parts.append(f"<p class='degree'>{html.escape(data['degree'])}</p>")
            parts.append("</section>")
        elif kind == "pagebreak":
            parts.append("<div class='pagebreak'></div>")
        elif kind == "front-h1":
            parts.append(f"<h1 class='front'>{html.escape(block[1])}</h1>")
        elif kind == "h1":
            parts.append(f"<h1>{inline_html(block[1])}</h1>")
        elif kind == "h2":
            parts.append(f"<h2>{inline_html(block[1])}</h2>")
        elif kind == "p":
            parts.append(f"<p>{inline_html(block[1])}</p>")
        elif kind == "bullets":
            parts.append("<ul>")
            for item in block[1]:
                parts.append(f"<li>{inline_html(item)}</li>")
            parts.append("</ul>")
        elif kind == "toc":
            parts.append("<div class='toc'>")
            for level, title, key in TOC:
                cls = "toc-row" + (" lvl2" if level == 2 else "")
                parts.append(
                    f"<div class='{cls}'><span class='t'>{html.escape(title)}</span>"
                    f"<span class='dots'></span><span class='pg'>{page_no(page_numbers, key)}</span></div>"
                )
            parts.append("</div>")
        elif kind == "figlist":
            parts.append("<table class='list-table'>")
            for fig in FIGURES:
                parts.append(
                    f"<tr><td>Figure {fig['no']}. {inline_html(fig['caption'])}</td>"
                    f"<td class='pg'>{page_no(page_numbers, fig['key'])}</td></tr>"
                )
            parts.append("</table>")
        elif kind == "tablist":
            parts.append("<table class='list-table'>")
            for tab in TABLES:
                parts.append(
                    f"<tr><td>Table {tab['no']}. {inline_html(tab['caption'])}</td>"
                    f"<td class='pg'>{page_no(page_numbers, tab['key'])}</td></tr>"
                )
            parts.append("</table>")
        elif kind == "abbreviations":
            parts.append("<table class='abbr'>")
            for abbrev, meaning in block[1]:
                parts.append(f"<tr><td style='width:16%'><strong>{html.escape(abbrev)}</strong></td><td>{html.escape(meaning)}</td></tr>")
            parts.append("</table>")
        elif kind == "table":
            number = next(t["no"] for t in TABLES if t["caption"] == block[1]["caption"])
            data = block[1]
            parts.append("<div class='table-block'>")
            parts.append(f"<div class='table-caption'><strong>Table {number}.</strong> {inline_html(data['caption'])}</div>")
            parts.append("<table class='data'><colgroup>")
            for width in data["widths"]:
                parts.append(f"<col style='width:{width * 100:.2f}%'>")
            parts.append("</colgroup><thead><tr>")
            for column in data["columns"]:
                parts.append(f"<th>{inline_html(column)}</th>")
            parts.append("</tr></thead><tbody>")
            for row in data["rows"]:
                parts.append("<tr>")
                for cell in row:
                    parts.append(f"<td>{inline_html(cell)}</td>")
                parts.append("</tr>")
            parts.append("</tbody></table></div>")
        elif kind == "figure":
            number = next(f["no"] for f in FIGURES if f["caption"] == block[1]["caption"])
            data = block[1]
            parts.append(f"<figure class='figure-block' id='fig-{number}'>")
            parts.append(f"<img src='{html.escape(data['file'])}' style='width:{data['width']}in' alt='Figure {number}'>")
            parts.append(f"<figcaption class='figure-caption'><strong>Figure {number}.</strong> {inline_html(data['caption'])}</figcaption>")
            parts.append("</figure>")
        elif kind == "listing":
            data = block[1]
            parts.append("<div class='code-block'>")
            parts.append(f"<div class='code-caption'>{inline_html(data['caption'])}</div>")
            parts.append(f"<pre>{html.escape(chr(10).join(data['lines']))}</pre>")
            parts.append("</div>")
        elif kind == "code":
            parts.append("<div class='code-block'>")
            parts.append(f"<pre>{html.escape(chr(10).join(block[1]['lines']))}</pre>")
            parts.append("</div>")
        elif kind == "refs":
            for number, text in block[1]:
                parts.append(f"<p class='ref'>{html.escape(number)} {inline_html(text)}</p>")
        elif kind == "qa":
            for question, answer in block[1]:
                parts.append(
                    f"<p class='qa'><span class='q'>{inline_html(question)}</span> "
                    f"{inline_html(answer)}</p>"
                )
        else:
            raise ValueError(f"unknown block {kind}")

    parts.append("</body></html>")
    return "\n".join(parts)


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_margins(cell, top=40, start=60, bottom=40, end=60) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def add_first_page_border(section) -> None:
    sect_pr = section._sectPr
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    borders.set(qn("w:display"), "firstPage")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:space"), "24")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    cols = sect_pr.find(qn("w:cols"))
    if cols is not None:
        cols.addprevious(borders)
    else:
        sect_pr.append(borders)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run._r.append(instr)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.append(separate)
    if placeholder:
        run.add_text(placeholder)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def set_paragraph_border(paragraph, position: str = "bottom") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    el = OxmlElement(f"w:{position}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), "4")
    el.set(qn("w:space"), "1")
    el.set(qn("w:color"), "8A96A3")
    borders.append(el)


def add_inline_runs(paragraph, text: str, bold: bool = False, size: float | None = None) -> None:
    token = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
    position = 0
    for match in token.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            run.bold = bold
            if size:
                run.font.size = Pt(size)
        value = match.group(0)
        if value.startswith("**"):
            run = paragraph.add_run(value[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(value[1:-1])
            run.font.name = FONT_MONO
            run.font.size = Pt((size or 11) - 1.2)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_MONO)
        if size:
            run.font.size = Pt(size)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        run.bold = bold
        if size:
            run.font.size = Pt(size)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.1)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = FONT_BODY
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, before, after in (
        ("Heading 1", 14, 12, 6),
        ("Heading 2", 12, 10, 4),
        ("Heading 3", 11, 8, 3),
    ):
        style = doc.styles[style_name]
        style.font.name = FONT_BODY
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    run = hp.add_run("DISTRIBUTED COMPUTING SYSTEM | MINI PROJECT REPORT")
    run.font.name = FONT_BODY
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(80, 80, 80)
    set_paragraph_border(hp, "bottom")

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_after = Pt(0)
    run = fp.add_run("RailReserve | Final Year B.Tech IT | Page ")
    run.font.name = FONT_BODY
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(80, 80, 80)
    add_field(fp, " PAGE ", "1")
    run = fp.add_run(" of ")
    run.font.name = FONT_BODY
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(80, 80, 80)
    add_field(fp, " NUMPAGES ", "1")

    first_footer = section.first_page_footer
    first_footer.is_linked_to_previous = False
    ffp = first_footer.paragraphs[0]
    ffp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ffp.paragraph_format.space_after = Pt(0)
    add_field(ffp, " PAGE ", "1")

    add_first_page_border(section)


def add_caption(doc: Document, text: str, before: bool = True) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4 if before else 3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = before
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9.4)
    run.font.name = FONT_BODY


def add_code_table(doc: Document, lines: list[str], caption: str | None = None) -> None:
    if caption:
        add_caption(doc, caption, before=True)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(CONTENT_WIDTH_IN)
    cell = table.rows[0].cells[0]
    cell.width = Inches(CONTENT_WIDTH_IN)
    set_cell_shading(cell, "F5F7FA")
    set_cell_margins(cell, top=70, start=110, bottom=70, end=110)
    set_cant_split(table.rows[0])
    cell.paragraphs[0].text = ""
    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(line if line else " ")
        run.font.name = FONT_MONO
        run.font.size = Pt(8.3)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_MONO)
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "C9D0D8")
        borders.append(el)
    tbl_pr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_data_table(doc: Document, data: dict, number: int) -> None:
    add_caption(doc, f"Table {number}. {data['caption']}", before=True)
    columns = data["columns"]
    rows = data["rows"]
    widths = data["widths"]
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(CONTENT_WIDTH_IN * width)
    header = table.rows[0]
    set_repeat_header(header)
    set_cant_split(header)
    for index, column in enumerate(columns):
        cell = header.cells[index]
        cell.width = Inches(CONTENT_WIDTH_IN * widths[index])
        set_cell_shading(cell, "DCE6F1")
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        add_inline_runs(paragraph, column, bold=True, size=8.8)
    for row in rows:
        cells = table.add_row().cells
        set_cant_split(table.rows[-1])
        for index, value in enumerate(row):
            cell = cells[index]
            cell.width = Inches(CONTENT_WIDTH_IN * widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            add_inline_runs(paragraph, value, size=8.8)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_figure(doc: Document, figure: dict) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(ASSETS / Path(figure["file"]).name), width=Inches(min(figure["width"], 6.3)))
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(7)
    run = caption.add_run(f"Figure {figure['no']}. {figure['caption']}")
    run.italic = True
    run.font.size = Pt(9.4)
    run.font.name = FONT_BODY


def add_toc_entry(doc: Document, title: str, key: str, level: int, page_numbers: dict) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    if level == 2:
        paragraph.paragraph_format.left_indent = Cm(0.7)
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Cm(16.8), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )
    run = paragraph.add_run(title)
    run.font.name = FONT_BODY
    run.font.size = Pt(9.4 if level == 1 else 9.1)
    if level == 1:
        run.bold = True
    run = paragraph.add_run("\t" + page_no(page_numbers, key))
    run.font.name = FONT_BODY
    run.font.size = Pt(9.2)


def add_list_table(doc: Document, items: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Inches(6.1)
    table.columns[1].width = Inches(0.4)
    for label, number in items:
        cells = table.add_row().cells
        cells[0].width = Inches(6.1)
        cells[1].width = Inches(0.4)
        p = cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        add_inline_runs(p, label, size=9.2)
        p = cells[1].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(number)
        run.font.size = Pt(9.2)
        run.font.name = FONT_BODY
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def docx_report(page_numbers: dict) -> Document:
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Preventing Double Booking in a Railway Reservation System"
    doc.core_properties.subject = "DCS FA2 Mini Project Report - RailReserve"
    doc.core_properties.author = "Tony Prajapati, Vedant Patarwale, Pratik Rajput, Ishan Toraskar"
    doc.core_properties.keywords = "distributed systems, concurrency, idempotency, railway reservation"

    for block in BLOCKS:
        kind = block[0]
        if kind == "title":
            data = block[1]
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(data["trust"])
            r.font.size = Pt(10.4)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(data["college"])
            r.bold = True
            r.font.size = Pt(11.2)
            for key in ("address", "autonomy", "affiliation"):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(data[key])
                r.font.size = Pt(10.2)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(data["department"])
            r.bold = True
            r.font.size = Pt(11.4)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(data["year"])
            r.font.size = Pt(10.8)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(6)
            p.add_run().add_picture(str(ASSETS / "pccoe-logo.png"), width=Inches(1.08))
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(data["report"])
            r.font.size = Pt(10.6)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(data["title"])
            r.bold = True
            r.font.size = Pt(18.5)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(18)
            r = p.add_run(data["subtitle"])
            r.italic = True
            r.font.size = Pt(10.8)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(data["by"])
            r.font.size = Pt(10.6)
            for name, roll in data["students"]:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(f"{name}: {roll}")
                r.font.size = Pt(11)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(16)
            r = p.add_run(data["degree"])
            r.bold = True
            r.font.size = Pt(10.8)
            doc.add_page_break()
        elif kind == "pagebreak":
            doc.add_page_break()
        elif kind == "front-h1":
            doc.add_heading(block[1], level=1)
        elif kind == "h1":
            doc.add_heading(block[1], level=1)
        elif kind == "h2":
            doc.add_heading(block[1], level=2)
        elif kind == "p":
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline_runs(paragraph, block[1])
        elif kind == "bullets":
            for item in block[1]:
                paragraph = doc.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.space_after = Pt(2.5)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                add_inline_runs(paragraph, item)
        elif kind == "toc":
            for level, title, key in TOC:
                add_toc_entry(doc, title, key, level, page_numbers)
        elif kind == "figlist":
            add_list_table(
                doc,
                [
                    (f"Figure {fig['no']}. {fig['caption']}", page_no(page_numbers, fig["key"]))
                    for fig in FIGURES
                ],
            )
        elif kind == "tablist":
            add_list_table(
                doc,
                [
                    (f"Table {tab['no']}. {tab['caption']}", page_no(page_numbers, tab["key"]))
                    for tab in TABLES
                ],
            )
        elif kind == "abbreviations":
            table = doc.add_table(rows=0, cols=2)
            table.style = "Table Grid"
            table.autofit = False
            table.columns[0].width = Inches(1.0)
            table.columns[1].width = Inches(5.5)
            for abbrev, meaning in block[1]:
                cells = table.add_row().cells
                cells[0].width = Inches(1.0)
                cells[1].width = Inches(5.5)
                set_cell_margins(cells[0], top=20, bottom=20)
                set_cell_margins(cells[1], top=20, bottom=20)
                p = cells[0].paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(abbrev)
                r.bold = True
                r.font.size = Pt(9.2)
                p = cells[1].paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(meaning)
                r.font.size = Pt(9.2)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif kind == "table":
            number = next(t["no"] for t in TABLES if t["caption"] == block[1]["caption"])
            add_data_table(doc, block[1], number)
        elif kind == "figure":
            figure = next(f for f in FIGURES if f["caption"] == block[1]["caption"])
            add_figure(doc, figure)
        elif kind == "listing":
            add_code_table(doc, block[1]["lines"], block[1]["caption"])
        elif kind == "code":
            add_code_table(doc, block[1]["lines"])
        elif kind == "refs":
            for number, text in block[1]:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.8)
                paragraph.paragraph_format.first_line_indent = Cm(-0.8)
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = paragraph.add_run(f"{number} {text}")
                run.font.size = Pt(9.3)
                run.font.name = FONT_BODY
        elif kind == "qa":
            for question, answer in block[1]:
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                paragraph.paragraph_format.space_after = Pt(5)
                run = paragraph.add_run(question + " ")
                run.bold = True
                add_inline_runs(paragraph, answer)
        else:
            raise ValueError(f"unknown block {kind}")

    return doc


def main() -> None:
    mode = sys.argv[1] if len(sys.argv) > 1 else "html"
    page_numbers_path = ROOT / "page_numbers.json"
    page_numbers = (
        json.loads(page_numbers_path.read_text()) if page_numbers_path.exists() else {}
    )
    if mode == "html":
        (ROOT / "report_draft.html").write_text(html_page({}))
        print(ROOT / "report_draft.html")
    elif mode == "html-final":
        (ROOT / "report_final.html").write_text(html_page(page_numbers))
        print(ROOT / "report_final.html")
    elif mode == "docx":
        doc = docx_report(page_numbers)
        output = REPO / "RailReserve_MiniProject_DCS.docx"
        doc.save(output)
        print(output)
    else:
        raise SystemExit(f"unknown mode: {mode}")


if __name__ == "__main__":
    main()
